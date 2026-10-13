import csv

from backend.ingestion.file_loader import extract_file_data


def test_html(tmp_path):
    p = tmp_path / "p.html"
    p.write_text(
        "<html><head><title>T</title><script>bad()</script></head>"
        "<body><h1>Hello</h1><table><tr><th>a</th><th>b</th></tr>"
        "<tr><td>1</td><td>2</td></tr></table></body></html>"
    )
    text = extract_file_data(str(p))["text"]
    assert "Hello" in text and "a | b" in text and "bad()" not in text


def test_docx(tmp_path):
    from docx import Document

    d = Document()
    d.add_paragraph("Quarterly report")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "x"
    t.rows[0].cells[1].text = "y"
    p = tmp_path / "r.docx"
    d.save(p)
    text = extract_file_data(str(p))["text"]
    assert "Quarterly report" in text and "x | y" in text


def test_xlsx(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.append(["name", "qty"])
    wb.active.append(["bolt", 5])
    p = tmp_path / "s.xlsx"
    wb.save(p)
    text = extract_file_data(str(p))["text"]
    assert "bolt | 5" in text


def test_pptx(tmp_path):
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Roadmap"
    slide.placeholders[1].text = "Ship v2"
    p = tmp_path / "d.pptx"
    prs.save(p)
    data = extract_file_data(str(p))
    assert "Roadmap" in data["text"] and "Ship v2" in data["text"]
    assert data["pages"] == 1


def test_csv_and_text_encodings(tmp_path):
    p = tmp_path / "c.csv"
    with open(p, "w", newline="", encoding="cp1252") as f:
        csv.writer(f).writerow(["café", "1"])
    assert "café" in extract_file_data(str(p))["text"]


def test_pdf(tmp_path):
    from pypdf import PdfWriter
    import pytest

    # blank page: loader must not crash (OCR fallback is best effort)
    w = PdfWriter()
    w.add_blank_page(200, 200)
    p = tmp_path / "b.pdf"
    with open(p, "wb") as f:
        w.write(f)
    assert extract_file_data(str(p))["pages"] == 1


def test_unsupported(tmp_path):
    import pytest

    p = tmp_path / "x.exe"
    p.write_bytes(b"MZ")
    with pytest.raises(ValueError):
        extract_file_data(str(p))


def test_scanned_pdf_uses_ocr(tmp_path):
    import pytest

    pytest.importorskip("rapidocr")
    from PIL import Image, ImageDraw, ImageFont

    font = ImageFont.truetype(
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 34
    )
    img = Image.new("RGB", (1240, 400), "white")
    ImageDraw.Draw(img).text(
        (40, 80), "Dividend distribution policy", fill="black", font=font
    )
    pdf = tmp_path / "scan.pdf"
    img.save(pdf)  # image-only PDF, no text layer

    text = extract_file_data(str(pdf))["text"]
    assert "Dividend" in text
