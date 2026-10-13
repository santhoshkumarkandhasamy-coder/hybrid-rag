"""Read text out of every supported upload format."""

import shutil
import subprocess
import tempfile
from pathlib import Path

# ============================================================
# SUPPORTED FILE TYPES
# ============================================================

# Legacy binary formats are converted with LibreOffice first.
LEGACY_CONVERSIONS = {
    ".doc": "docx",
    ".xls": "xlsx",
    ".ppt": "pptx",
}

IMAGE_EXTENSIONS = {
    ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp",
}

HTML_EXTENSIONS = {".html", ".htm"}

TEXT_EXTENSIONS = {".txt", ".md", ".csv"}

SUPPORTED_EXTENSIONS = (
    {".pdf", ".docx", ".xlsx", ".pptx"}
    | set(LEGACY_CONVERSIONS)
    | IMAGE_EXTENSIONS
    | HTML_EXTENSIONS
    | TEXT_EXTENSIONS
)

# Minimum characters on a PDF page before we consider it "scanned"
_PDF_OCR_THRESHOLD = 20


# ============================================================
# LIBREOFFICE (legacy .doc / .xls / .ppt)
# ============================================================

def find_libreoffice() -> str:

    soffice = shutil.which("soffice") or shutil.which("libreoffice")

    if soffice:
        return soffice

    for path in (
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        "/usr/bin/soffice",
        "/usr/local/bin/soffice",
    ):
        if Path(path).exists():
            return path

    raise RuntimeError(
        "LibreOffice is not installed or could not be found."
    )


def convert_with_libreoffice(file_path: str, target: str) -> str:
    """Convert a legacy Office file to `target` (docx/xlsx/pptx)."""

    input_path = Path(file_path)

    output_dir = Path(tempfile.mkdtemp(prefix="office_conversion_"))

    try:
        result = subprocess.run(
            [
                find_libreoffice(),
                "--headless",
                "--convert-to", target,
                "--outdir", str(output_dir),
                str(input_path),
            ],
            capture_output=True,
            text=True,
            timeout=180,
        )
    except subprocess.TimeoutExpired:
        raise RuntimeError("LibreOffice conversion timed out.")

    converted = output_dir / f"{input_path.stem}.{target}"

    if result.returncode != 0 or not converted.exists():
        raise RuntimeError(
            f"Could not convert {input_path.suffix} file: "
            f"{result.stderr.strip()}"
        )

    return str(converted)


# ============================================================
# OCR
# ============================================================

def _ocr_image(image) -> str:

    import pytesseract

    if image.mode not in ("RGB", "L"):
        image = image.convert("RGB")

    # Upscale small images; tesseract does poorly below ~300 DPI.
    if max(image.size) < 1500:
        scale = 1500 / max(image.size)
        image = image.resize(
            (int(image.width * scale), int(image.height * scale))
        )

    return pytesseract.image_to_string(image)


# ============================================================
# PDF
# ============================================================

def extract_pdf(file_path: str) -> dict:

    from pypdf import PdfReader

    reader = PdfReader(file_path)

    pages = []

    for number, page in enumerate(reader.pages, start=1):

        text = (page.extract_text() or "").strip()

        # Scanned page: fall back to OCR if tooling is available.
        if len(text) < _PDF_OCR_THRESHOLD:
            text = _ocr_pdf_page(file_path, number - 1) or text

        if text:
            pages.append(f"[Page {number}]\n{text}")

    return {
        "file_type": "pdf",
        "pages": len(reader.pages),
        "text": "\n\n".join(pages),
    }


def _ocr_pdf_page(file_path: str, index: int) -> str:

    try:
        import pypdfium2

        pdf = pypdfium2.PdfDocument(file_path)
        image = pdf[index].render(scale=2.5).to_pil()

        return _ocr_image(image).strip()

    except Exception:
        return ""


# ============================================================
# WORD
# ============================================================

def extract_docx(file_path: str) -> dict:

    from docx import Document

    document = Document(file_path)

    sections = []

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            sections.append(text)

    for index, table in enumerate(document.tables, start=1):

        sections.append(f"\n[Table {index}]")

        for row in table.rows:
            sections.append(
                " | ".join(cell.text.strip() for cell in row.cells)
            )

    return {
        "file_type": "docx",
        "pages": None,
        "text": "\n".join(sections),
    }


# ============================================================
# EXCEL
# ============================================================

def extract_xlsx(file_path: str) -> dict:

    from openpyxl import load_workbook

    workbook = load_workbook(
        filename=file_path, data_only=True, read_only=True
    )

    sections = []

    for worksheet in workbook.worksheets:

        sections.append(f"\n[Worksheet: {worksheet.title}]")

        for row in worksheet.iter_rows(values_only=True):

            values = [str(v) for v in row if v is not None]

            if values:
                sections.append(" | ".join(values))

    workbook.close()

    return {
        "file_type": "xlsx",
        "pages": len(workbook.worksheets),
        "text": "\n".join(sections),
    }


# ============================================================
# POWERPOINT
# ============================================================

def extract_pptx(file_path: str) -> dict:

    from pptx import Presentation

    presentation = Presentation(file_path)

    slides = []

    for number, slide in enumerate(presentation.slides, start=1):

        parts = [f"[Slide {number}]"]

        for shape in slide.shapes:

            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    parts.append(text)

            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(
                        " | ".join(
                            cell.text.strip() for cell in row.cells
                        )
                    )

        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"Notes: {notes}")

        if len(parts) > 1:
            slides.append("\n".join(parts))

    return {
        "file_type": "pptx",
        "pages": len(presentation.slides),
        "text": "\n\n".join(slides),
    }


# ============================================================
# HTML
# ============================================================

def extract_html(file_path: str) -> dict:

    from bs4 import BeautifulSoup

    raw = Path(file_path).read_bytes()

    soup = BeautifulSoup(raw, "html.parser")

    for tag in soup(["script", "style", "noscript", "template"]):
        tag.decompose()

    for table in soup.find_all("table"):

        rows = [
            " | ".join(
                cell.get_text(" ", strip=True)
                for cell in row.find_all(["th", "td"])
            )
            for row in table.find_all("tr")
        ]

        table.replace_with("\n[Table]\n" + "\n".join(rows) + "\n")

    title = soup.title.get_text(strip=True) if soup.title else ""

    body = soup.get_text("\n", strip=True)

    text = f"{title}\n\n{body}" if title else body

    return {"file_type": "html", "pages": None, "text": text}


# ============================================================
# PLAIN TEXT / CSV
# ============================================================

def extract_text(file_path: str) -> dict:

    raw = Path(file_path).read_bytes()

    # utf-16 only when a BOM says so (it otherwise "decodes" anything)
    encodings = ["utf-8-sig", "cp1252", "latin-1"]

    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        encodings.insert(0, "utf-16")

    for encoding in encodings:
        try:
            text = raw.decode(encoding)
            break
        except UnicodeError:
            continue

    suffix = Path(file_path).suffix.lower().lstrip(".")

    return {"file_type": suffix, "pages": None, "text": text}


# ============================================================
# IMAGES
# ============================================================

def extract_image(file_path: str) -> dict:

    from PIL import Image

    with Image.open(file_path) as image:
        text = _ocr_image(image)

    return {"file_type": "image", "pages": None, "text": text}


# ============================================================
# UNIVERSAL LOADER
# ============================================================

def extract_file_data(file_path: str) -> dict:

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    extension = path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type: {extension}. Supported: "
            + ", ".join(sorted(SUPPORTED_EXTENSIONS))
        )

    if extension in LEGACY_CONVERSIONS:

        target = LEGACY_CONVERSIONS[extension]
        converted = convert_with_libreoffice(file_path, target)

        try:
            data = globals()[f"extract_{target}"](converted)
        finally:
            shutil.rmtree(Path(converted).parent, ignore_errors=True)

        data["file_type"] = extension.lstrip(".")
        return data

    if extension == ".pdf":
        return extract_pdf(file_path)

    if extension == ".docx":
        return extract_docx(file_path)

    if extension == ".xlsx":
        return extract_xlsx(file_path)

    if extension == ".pptx":
        return extract_pptx(file_path)

    if extension in HTML_EXTENSIONS:
        return extract_html(file_path)

    if extension in TEXT_EXTENSIONS:
        return extract_text(file_path)

    return extract_image(file_path)
