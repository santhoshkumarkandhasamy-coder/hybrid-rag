import openpyxl,sys,warnings
warnings.filterwarnings('ignore')
"""Read-only dump (formula + cached value) of a large reference workbook.
Usage: python3 -I dump_reference.py <file.xlsx> <sheet> <row_from> <row_to> [max_col]"""
f=sys.argv[1]; s=sys.argv[2]; r0=int(sys.argv[3]); r1=int(sys.argv[4]); cmax=int(sys.argv[5]) if len(sys.argv)>5 else 14
wf=openpyxl.load_workbook(f,read_only=True)[s]
wv=openpyxl.load_workbook(f,read_only=True,data_only=True)[s]
vals={}
for row in wv.iter_rows(min_row=r0,max_row=r1,max_col=cmax):
    for c in row:
        if getattr(c,'value',None) is not None: vals[c.coordinate]=c.value
for row in wf.iter_rows(min_row=r0,max_row=r1,max_col=cmax):
    out=[]
    for c in row:
        v=getattr(c,'value',None)
        if v is None: continue
        if isinstance(v,str) and v.startswith('='):
            vv=vals.get(c.coordinate); vv=round(vv,2) if isinstance(vv,float) else vv
            out.append(f"{c.coordinate}={str(vv)[:30]}{{{v[:40]}}}")
        else:
            v=round(v,2) if isinstance(v,float) else v
            out.append(f"{c.coordinate}:{str(v).replace(chr(10),' ')[:90]}")
    if out: print(' | '.join(out))
