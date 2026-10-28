"""Shared helpers for the stage-7 rebuild (reference-format notes)."""
import copy
from openpyxl.styles import Font, Border, Side, Alignment, PatternFill
from openpyxl.cell.cell import MergedCell

TNR = 'Times New Roman'
NUM = '_(* #,##0.00_);_(* (#,##0.00);_(* "-"??_);_(@_)'
INT = '_(* #,##0_);_(* (#,##0);_(* "-"??_);_(@_)'
PCT = '0.00%;(0.00%);"-"'
RAT = '0.00;(0.00);"-"'
THIN = Side(style='thin')
NOS = Side(style=None)

MD = 'MasterData'
DATE = {'D': '$B$9', 'E': '$B$11', 'F': '$B$12', 'G': '$B$13', 'H': '$B$14'}
PREV = {'D': '$B$11', 'E': '$B$12', 'F': '$B$13', 'G': '$B$14'}
TB = {'D': 'TB 2026-27', 'E': 'TB 2025-26', 'F': 'TB 2024-25', 'G': 'TB 2023-24'}
SC = 'MasterData!$B$15'
ADJ = "'Adj Entry RFS'!"


def as_at(col):
    return f'="As at \n"&TEXT(MasterData!{DATE[col]},"mmmm dd, yyyy")'


def for_per(col):
    k = 'period' if col == 'D' else 'year'
    return f'="For the {k} ended\n"&TEXT(MasterData!{DATE[col]},"mmmm dd, yyyy")'


def tb(col, grp, ledger=None):
    t = TB[col]
    if ledger is None:
        return f"SUMIF('{t}'!$G:$G,\"{grp}\",'{t}'!$I:$I)"
    return f"SUMIFS('{t}'!$I:$I,'{t}'!$G:$G,\"{grp}\",'{t}'!$A:$A,\"{ledger}\")"


def adj_win(col, master, dr_positive=True, mcol='C'):
    """Dr - Cr of restatement entries for `master` in the period of `col`."""
    lo, hi = PREV[col], DATE[col]
    def s(c):
        return (f"SUMIFS({ADJ}${c}$3:${c}$150,{ADJ}${mcol}$3:${mcol}$150,\"{master}\","
                f"{ADJ}$A$3:$A$150,\">\"&MasterData!{lo},{ADJ}$A$3:$A$150,\"<=\"&MasterData!{hi})")
    return f"({s('E')}-{s('F')})" if dr_positive else f"({s('F')}-{s('E')})"


def adj_cum(col, master):
    """Cumulative Dr - Cr of restatement entries for `master` up to the date of `col`."""
    hi = DATE[col]
    def s(c):
        return (f"SUMIFS({ADJ}${c}$3:${c}$150,{ADJ}$C$3:$C$150,\"{master}\","
                f"{ADJ}$A$3:$A$150,\"<=\"&MasterData!{hi})")
    return f"({s('E')}-{s('F')})"


def clear(ws, r0, r1=None, c0=1, c1=30):
    r1 = r1 or ws.max_row
    for m in list(ws.merged_cells.ranges):
        if m.max_row >= r0 and m.min_row <= r1:
            ws.unmerge_cells(str(m))
    for r in range(r0, r1 + 1):
        for c in range(c0, c1 + 1):
            x = ws.cell(r, c)
            x.value = None
            x.font = Font(name=TNR, size=10)
            x.border = Border()
            x.number_format = 'General'
            x.alignment = Alignment()
            x.fill = PatternFill()
        ws.row_dimensions[r].height = None


def put(ws, addr, v, b=False, i=False, nf=None, al=None, wrap=False, bd=None, sz=10, valign=None, color=None):
    x = ws[addr]
    if isinstance(x, MergedCell):
        return
    x.value = v
    x.font = Font(name=TNR, size=sz, bold=b, italic=i, color=color)
    if nf:
        x.number_format = nf
    elif isinstance(v, (int, float)) or (isinstance(v, str) and v.startswith('=') and not v.startswith('="')):
        x.number_format = NUM
    x.alignment = Alignment(horizontal=al, vertical=valign or ('center' if al == 'center' else 'top' if wrap else None), wrap_text=wrap)
    if bd:
        x.border = border(bd)
    return x


def border(spec):
    return Border(left=THIN if 'l' in spec else NOS, right=THIN if 'r' in spec else NOS,
                  top=THIN if 't' in spec else NOS, bottom=THIN if 'b' in spec else NOS)


def add_border(ws, addr, spec):
    x = ws[addr]
    o = x.border
    x.border = Border(left=THIN if 'l' in spec else o.left, right=THIN if 'r' in spec else o.right,
                      top=THIN if 't' in spec else o.top, bottom=THIN if 'b' in spec else o.bottom)


def box(ws, r0, r1, c0, c1, inner_v=True, inner_h=False):
    """Grid like the reference ageing / MSME tables."""
    from openpyxl.utils import get_column_letter as L
    for r in range(r0, r1 + 1):
        for c in range(c0, c1 + 1):
            s = ''
            if c == c0 or inner_v: s += 'l'
            if c == c1 or inner_v: s += 'r'
            if r == r0 or inner_h: s += 't'
            if r == r1 or inner_h: s += 'b'
            add_border(ws, f'{L(c)}{r}', s)


def hline(ws, r, cols, spec='tb'):
    for c in cols:
        add_border(ws, f'{c}{r}', spec)


def text(ws, r, v, c0='C', c1='H', b=False, i=False, h=None, al='justify'):
    put(ws, f'{c0}{r}', v, b=b, i=i, wrap=True, al=al)
    if c1 and c1 != c0:
        ws.merge_cells(f'{c0}{r}:{c1}{r}')
    if h is None and isinstance(v, str):
        h = max(14, 13.5 * (len(v) // 120 + 1))
    if h:
        ws.row_dimensions[r].height = h


def find(ws, label, col='C', start=1, exact=False):
    for r in range(start, ws.max_row + 1):
        v = ws[f'{col}{r}'].value
        if isinstance(v, str) and ((v.strip() == label) if exact else (label in v)):
            return r
    raise KeyError(f'{ws.title}: {label}')
